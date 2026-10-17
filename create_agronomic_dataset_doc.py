from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("agronomic_dataset_characteristics.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="808080", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(8)
heading.paragraph_format.keep_with_next = True
run = heading.add_run("3.3.2 Dataset")
set_run_font(run, size=12, bold=True)

add_body_paragraph(
    doc,
    "The agronomic-input dataset was developed to support the VISCANE sugarcane yield-estimation and recommendation component. It represents structured farm-management information supplied by the farmer after image analysis. The records combine the selected sugarcane variety, cultivated area, field preparation practices, crop-cycle stage, and RSSI infection status with the corresponding predicted production values. These inputs allow the system to estimate LKG/TC, tons cane per hectare, and total LKG while generating recommendations related to fertilizer application, weeding, plowing, and crop management."
)

add_body_paragraph(
    doc,
    "The dataset contained 24 records distributed equally among three sugarcane varieties: VMC 84-524, VMC 84-947, and MAURITIO RC888. The management variables were encoded as numeric values, generally using levels from 1 to 3, while cultivated area was recorded in hectares. The data was organized in comma-separated values format and was used as a structured reference for the agronomic calculation workflow. Because the dataset is small and represents system-generated or demonstration values, it should be interpreted as a prototype decision-support dataset rather than as a statistically representative field trial."
)

caption = doc.add_paragraph()
caption.paragraph_format.space_before = Pt(4)
caption.paragraph_format.space_after = Pt(4)
caption.paragraph_format.keep_with_next = True
caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = caption.add_run("Table 3.3. Characteristics of the Agronomic-Input Dataset")
set_run_font(run, size=11, bold=True, italic=True)

rows = [
    ("Dataset source", "Structured agronomic records stored in data/agronomic_training.csv for the VISCANE estimation workflow."),
    ("Total number of records", "24 agronomic-input records."),
    ("Sugarcane varieties", "VMC 84-524, VMC 84-947, and MAURITIO RC888; 8 records per variety."),
    ("Agronomic input variables", "Cultivated area, plowing count, weeding count, fertilizer count, ratoon stage, and RSSI infection status."),
    ("Cultivated area", "Recorded in hectares, ranging from 1.0 to 4.0 hectares."),
    ("Management-practice scale", "Plowing, weeding, and fertilizer counts were represented using levels from 1 to 3."),
    ("Crop-cycle variable", "Ratoon stage values from 1 to 3, representing New Plant, 1st ratoon, and 2nd ratoon."),
    ("Disease or infection variable", "RSSI infection status was stored as numeric values from 1 to 3 in the CSV dataset."),
    ("Prediction targets", "Predicted LKG/TC, predicted TC/HA, and predicted total LKG."),
    ("Target value ranges", "LKG/TC: 1.46–2.24; TC/HA: 242–690; total LKG: 384.80–1,367.28."),
    ("File format", "Comma-separated values (CSV), with one header row and ten fields per record."),
    ("Primary use", "Prototype yield estimation, agronomic recommendations, and agronomic-log recording."),
]

table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.allow_autofit = False
table.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_table_borders(table)
header = table.rows[0]
set_repeat_table_header(header)
for i, text in enumerate(("Dataset Characteristic", "Description")):
    set_cell_width(header.cells[i], (2.05, 4.45)[i])
    set_cell_margins(header.cells[i])
    set_cell_shading(header.cells[i], "D9E2F3")
    add_cell_text(header.cells[i], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

for characteristic, description in rows:
    cells = table.add_row().cells
    for i, width in enumerate((2.05, 4.45)):
        set_cell_width(cells[i], width)
        set_cell_margins(cells[i])
    add_cell_text(cells[0], characteristic, bold=True)
    add_cell_text(cells[1], description)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after = Pt(0)
note.paragraph_format.line_spacing = 1.15
run = note.add_run("Note. ")
set_run_font(run, size=10, bold=True)
run = note.add_run("The agronomic-input CSV contains 24 demonstration records. The running service applies variety-specific weighted baselines; therefore, the table describes the dataset used for the agronomic workflow and not a validated field-trial sample.")
set_run_font(run, size=10, italic=True)

doc.save(OUT)
print(OUT)
